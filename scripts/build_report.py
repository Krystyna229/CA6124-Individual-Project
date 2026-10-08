#!/usr/bin/env python3
"""Build the four-page CA6124 report from recorded experiment artifacts."""

import json
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report"


def margins(section):
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(.8)
    section.bottom_margin = Inches(.8)
    section.left_margin = Inches(.75)
    section.right_margin = Inches(.75)


def shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def add_table(doc, columns, rows, widths=None, font_size=11):
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for index, title in enumerate(columns):
        cell = table.rows[0].cells[index]
        cell.text = title
        shade(cell, "D9EAF7")
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            if widths:
                cell.width = Inches(widths[index])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    if row is table.rows[0]:
                        run.bold = True
    return table


def add_hyperlink(paragraph, text, url):
    relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), "Arial")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([fonts, color, underline])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def set_style_font(style, name, size=None, bold=None):
    style.font.name = name
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def make_ranking_curves():
    ks = [1, 5, 10, 20, 30, 50]
    models = ["EASE", "Multi-VAE", "CDAE"]
    metrics = ["NDCG", "Hit Ratio", "Precision"]
    datasets = [("ml-1m", "ML-1M"), ("amazon-music", "Amazon Music")]
    colors = {"EASE": "#4C78A8", "Multi-VAE": "#F58518", "CDAE": "#54A24B"}
    markers = {"EASE": "o", "Multi-VAE": "s", "CDAE": "^"}
    fig, axes = plt.subplots(2, 3, figsize=(7.25, 3.35), sharex=True)
    for row, (dataset, label) in enumerate(datasets):
        for col, metric in enumerate(metrics):
            ax = axes[row, col]
            for model in models:
                slug = model.lower().replace("multi-vae", "multi-vae")
                frame = pd.read_csv(ROOT / f"results/kpi/{dataset}_{slug}.csv", index_col=0)
                values = frame.loc[metric].astype(float).to_numpy()
                ax.plot(ks, values, marker=markers[model], linewidth=1.4,
                        markersize=3.5, color=colors[model], label=model)
            ax.grid(axis="y", color="#D9D9D9", linewidth=.55)
            ax.tick_params(labelsize=7)
            ax.set_xticks(ks)
            if row == 0:
                ax.set_title(metric, fontsize=8.5, fontweight="bold")
            if col == 0:
                ax.set_ylabel(f"{label}\nscore", fontsize=8)
            if row == 1:
                ax.set_xlabel("K", fontsize=8)
    handles, labels = axes[0, 0].get_legend_handles_labels()
    fig.legend(handles, labels, loc="upper center", ncol=3, frameon=False,
               fontsize=8, bbox_to_anchor=(.5, 1.01))
    fig.tight_layout(rect=[0, 0, 1, .94], pad=.8, w_pad=.8, h_pad=.7)
    path = ROOT / "results/ranking_curves.png"
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return path


def heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.add_run(text)
    return paragraph


def body(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.0
    return paragraph


def page_break(doc):
    doc.add_page_break()


def main():
    OUT.mkdir(exist_ok=True)
    diagnostics = json.loads((ROOT / "diagnostics/dataset_diagnostics.json").read_text())
    results = pd.read_csv(ROOT / "results/final_summary.csv")
    params = pd.read_csv(ROOT / "results/best_parameters.csv").fillna("")
    ranking_curves = make_ranking_curves()

    doc = Document()
    margins(doc.sections[0])
    styles = doc.styles
    set_style_font(styles["Normal"], "Arial", 11)
    for name, size in (("Heading 1", 13), ("Heading 2", 11)):
        set_style_font(styles[name], "Arial", size, bold=True)
        styles[name].font.color.rgb = __import__("docx").shared.RGBColor.from_string("000000")
        styles[name].paragraph_format.space_before = Pt(4)
        styles[name].paragraph_format.space_after = Pt(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Extending DaisyRec 2.0 with CDAE")
    run.bold = True
    run.font.size = Pt(16)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Warm-start full-ranking comparison on MovieLens 1M and Amazon Digital Music").italic = True
    ident = doc.add_paragraph()
    ident.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ident.add_run("Name: Yanyu Wen     Student ID: [YOUR STUDENT ID]").bold = True

    heading(doc, "1 Introduction")
    body(doc, "This project integrates Collaborative Denoising Auto-Encoder (CDAE) [1] into DaisyRec 2.0 and compares it with its EASE [2] and Multi-VAE [3] baselines under one reproducible warm-start full-ranking protocol. MovieLens 1M and the 2014 Amazon Digital Music ratings-only release provide dense movie and sparse e-commerce settings. Because they also differ in domain, scale and temporal coverage, results are interpreted as dataset-specific trends rather than causal effects of sparsity.")

    heading(doc, "1.1 Data diagnosis and fixed decision rule", 2)
    drows = []
    for label, key in (("ML-1M", "ml-1m"), ("Amazon Music", "amazon-music")):
        d = next(item for item in diagnostics[key]["variants"]
                 if item["preprocessing"] == "10filter")
        drows.append([label, f"{d['users']:,}", f"{d['items']:,}",
                      f"{d['interactions']:,}", f"{100*d['density']:.3f}%",
                      f"{d['actual_train_ratio']:.3f}",
                      f"{d['warm_test_interactions']:,}",
                      f"{d['warm_test_users']:,}"])
    add_table(doc, ["Dataset", "Users", "Items", "Events", "Density", "Train\nratio", "Warm test\ninteractions", "Eval.\nusers"], drows,
              widths=[1.05,.55,.6,.75,.65,.55,.9,.6], font_size=9.5)
    body(doc, "The rule was fixed before modelling: retain DaisyRec's one-pass 10filter unless fewer than 500 warm-start test users remain. Both datasets pass. Amazon 5filter was rejected because 20,356 items require a 3.09 GiB float64 item matrix before EASE's inverse/storage overhead. The original Amazon file has 836,006 headerless rows; its SHA-256 is recorded in README.")
    repository = doc.add_paragraph("Repository: ")
    repository.paragraph_format.space_after = Pt(3)
    add_hyperlink(repository, "github.com/Krystyna229/CA6124-Individual-Project",
                  "https://github.com/Krystyna229/CA6124-Individual-Project")

    page_break(doc)
    heading(doc, "2 CDAE and non-destructive DaisyRec extension")
    body(doc, "For binary history xᵤ, CDAE applies training-only dropout and computes hᵤ = ReLU(W₁x̃ᵤ + vᵤ + b₁), where vᵤ is a user embedding; decoder logits are zᵤ = W₂hᵤ + b₂. Training uses full-history binary cross-entropy with optional L1/L2 penalties, while eval() disables corruption. ReLU and full-history BCE, rather than the original paper's sigmoid/tanh and negative-sampling variants [1], are deliberate DaisyRec integration choices.")
    body(doc, "CDAE subclasses AERecommender and reuses DaisyRec's data, optimizer, device and training utilities. The added evaluation options - day split, warm start, full ranking, NDCG@10 tuning and per-user metrics - are opt-in; defaults and existing baselines remain unchanged. Seven regression tests cover the original defaults and all new paths.")

    heading(doc, "3 Experimental protocol")
    body(doc, "Ratings become implicit interactions after duplicate user-item removal. A deterministic boundary between complete calendar days is chosen nearest 80/20, keeping equal-date Amazon events together; Table 1 reports the resulting ratios.")
    body(doc, "Warm-start evaluation retains test users/items seen in training; tuning applies the same rule to its smaller training subset. Full ranking scores every training-seen item and masks prior positives. Final cohorts contain 1,152 ML-1M users / 105,699 interactions and 1,294 Amazon users / 4,065 interactions.")
    body(doc, "DaisyRec's tune.py (Optuna TPE, seed 2022) runs 20 trials per group on validation NDCG@10. Winners are retrained and tested once at K ∈ {1,5,10,20,30,50}. CDAE and Multi-VAE use Adam for 50 epochs without early stopping; their losses are BCE+L1/L2 and multinomial+KL, respectively. Multi-VAE anneals per mini-batch over 400 steps, so every searched cap is reached. EASE is closed-form.")

    heading(doc, "3.1 Search spaces", 2)
    add_table(doc, ["Model", "Search space"], [
        ["CDAE", "d {16,32,64,128,256}; drop {0,.1,.3,.5,.7}; batch {64,128,256,512}; lr {5e-5,1e-4,5e-4,.001,.005,.01}; L2 {0,1e-5,1e-4,.001,.01,.1,1}"],
        ["EASE", "regularization {1,5,10,50,100,200,500,1000,2000}"],
        ["Multi-VAE", "d {32,64,128,256,512}; drop {0,.2,.5,.7,.8}; batch {64,128,256,512}; lr {5e-5,1e-4,5e-4,.001,.005,.01}; anneal {.05,.1,.2,.5}"],
    ], widths=[1.05,5.85], font_size=10)

    page_break(doc)
    heading(doc, "4 Hyperparameter tuning")
    doc.add_picture(str(ROOT / "results/tuning_curves.png"), width=Inches(6.05))
    caption = doc.add_paragraph("Figure 1. Individual trial scores and best-so-far validation NDCG@10 (20 trials per group).")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(10)

    columns = [c for c in params.columns if c not in ("dataset", "model")]
    prows = []
    for _, row in params.iterrows():
        values = []
        for c in columns:
            if row[c] == "" or c == "ndcg":
                continue
            value = row[c]
            if c in ("latent_dim", "batch_size"):
                value = str(int(float(value)))
            elif isinstance(value, (float, int)):
                value = f"{float(value):g}"
            short = {"latent_dim": "d", "batch_size": "batch", "dropout": "drop",
                     "anneal_cap": "anneal", "reg_1": "L1", "reg_2": "L2"}.get(c, c)
            values.append(f"{short}={value}")
        selected = ", ".join(values)
        prows.append([row.dataset, row.model, selected, f"{float(row.ndcg):.4f}"])
    add_table(doc, ["Dataset", "Model", "Selected hyperparameters", "Val NDCG@10"], prows,
              widths=[.95,.8,4.25,.9], font_size=9.5)
    body(doc, "Amazon retains upper-bound optima (Multi-VAE d=512, anneal=0.5; CDAE d=256), so they are best tested settings rather than proven global optima; ML-1M Multi-VAE selects interior d=256 and anneal=0.1. Amazon is also less stable: 8/20 Multi-VAE and 10/20 CDAE trials fall below 0.01 NDCG@10 (4 and 8 below 0.005), while EASE remains within 0.0276-0.0358 and no ML-1M trial falls below 0.01.")

    heading(doc, "5 Final test results")
    rrows = []
    for _, row in results.iterrows():
        rrows.append([row.dataset, row.model, f"{row['recall@10']:.4f}", f"{row['ndcg@10']:.4f}"])
    add_table(doc, ["Dataset", "Model", "Recall@10", "NDCG@10"], rrows,
              widths=[1.4,1.3,1.1,1.1], font_size=10)
    doc.add_picture(str(ranking_curves), width=Inches(6.65))
    caption = doc.add_paragraph("Figure 2. NDCG, Hit Ratio and Precision across ranking cut-offs under the same full-ranking protocol.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(10)

    heading(doc, "6 Analysis and limitations")
    body(doc, "On ML-1M, CDAE has the highest NDCG@10 (0.4992), but Multi-VAE leads Recall@10 (0.0459 versus 0.0433). Their paired NDCG difference is only 0.00120 and its 10,000-resample 95% CI [-0.0111, 0.0136] crosses zero; both neural models reliably exceed EASE. Metric choice therefore changes the top model, while the nearly overlapping Hit Ratio curves offer little separation.")
    body(doc, "On Amazon, EASE leads NDCG@10 (0.0415), followed by CDAE (0.0400) and Multi-VAE (0.0383). Yet all paired 95% intervals cross zero: EASE-CDAE [-0.0072, 0.0100], EASE-Multi-VAE [-0.0050, 0.0117] and CDAE-Multi-VAE [-0.0074, 0.0109]. Thus no pairwise winner is reliable on this cohort.")
    body(doc, "Amazon averages 11.4 interactions per user versus 165.3 on ML-1M; short histories may weaken KL-regularized Multi-VAE and increase neural tuning sensitivity, while EASE uses direct co-occurrence. This is explanatory rather than causal because the datasets differ on several dimensions. Limitations are one split/seed, 20 trials, boundary optima and warm-start scope: only 1,152/6,040 (19.1%) ML-1M and 1,294/5,729 (22.6%) Amazon users are evaluated. Full ranking avoids sampled-candidate bias [7], but cold-start performance remains unknown.")

    heading(doc, "References", 2)
    refs = [
        "[1] Y. Wu et al., ‘Collaborative Denoising Auto-Encoders for Top-N Recommender Systems,’ WSDM, 2016.",
        "[2] H. Steck, ‘Embarrassingly Shallow Autoencoders for Sparse Data,’ WWW, 2019.",
        "[3] D. Liang et al., ‘Variational Autoencoders for Collaborative Filtering,’ WWW, 2018.",
        "[4] Z. Sun et al., ‘DaisyRec 2.0: Benchmarking Recommendation for Rigorous Evaluation,’ TPAMI, 2023.",
        "[5] F. Harper and J. Konstan, ‘The MovieLens Datasets: History and Context,’ TiiS, 2015.",
        "[6] J. McAuley et al., ‘Image-based Recommendations on Styles and Substitutes,’ SIGIR, 2015.",
        "[7] W. Krichene and S. Rendle, ‘On Sampled Metrics for Item Recommendation,’ KDD, 2020.",
    ]
    for reference in refs:
        paragraph = doc.add_paragraph(reference)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.left_indent = Inches(.15)
        paragraph.paragraph_format.first_line_indent = Inches(-.15)
        for run in paragraph.runs:
            run.font.size = Pt(10)

    for section in doc.sections:
        margins(section)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("CA6124 Individual Assignment")

    # Apply Arial explicitly to every textual run; this prevents Word theme
    # fonts from overriding the assignment's typography requirement.
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.name = "Arial"
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{key}"), "Arial")

    path = OUT / "CA6124_CDAE_Technical_Report.docx"
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
